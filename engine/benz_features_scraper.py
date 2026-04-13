"""
Mercedes-Benz Key Features scraper (mbusa.com individual vehicle pages).

확인된 DOM 구조:
  URL: https://www.mbusa.com/en/vehicles/model/{class}/{body}/{trim}

  Key Features 탭 활성화:
    a[href="#key-features"].sticky-nav__item-link  → JS click 필요
    (usercentrics-root overlay가 Playwright click 차단 → page.evaluate() 사용)

  피처 데이터 (Key Features 탭 클릭 후 DOM에 삽입):
    .model-page-features__options-list__container   × N카테고리
      └── p.model-page-features__options-list__title  → "Standard Features" | "Optional Features"
      └── ul
            └── li.model-page-features__options-list__item
                  └── button[data-category-name="Performance"][data-option-name="..."]

  카테고리는 button[data-category-name] 속성에 직접 기재됨.
  탭/카테고리 버튼 클릭 없이 전체 DOM 일괄 파싱 가능.

수집 전략:
  1. JS로 Key Features 탭 활성화 (overlay 우회)
  2. p.model-page-features__options-list__title 등장 대기
  3. 전체 DOM에서 카테고리·Standard/Optional 한 번에 추출
  4. 전체 트림 비교 후 Unavailable 처리

출력: CSV  Category, Feature, Trim1, Trim2, ...
      값:  Standard | Optional | Unavailable
"""

import csv
import time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError


CATEGORIES = ["Performance", "Design", "Safety", "Luxury", "Multimedia", "Accessories"]

# usercentrics 오버레이 처리 JS
_CONSENT_JS = """
() => {
    const root = document.getElementById('usercentrics-root');
    if (!root) return 'no-overlay';
    if (root.shadowRoot) {
        const btns = [...root.shadowRoot.querySelectorAll('button')];
        const accept = btns.find(b => /accept all|accept/i.test(b.textContent));
        if (accept) { accept.click(); return 'shadow-clicked'; }
    }
    root.style.display = 'none';
    root.style.pointerEvents = 'none';
    return 'overlay-hidden';
}
"""

# Key Features 탭 활성화 JS
_ACTIVATE_KF_JS = """
() => {
    const a = document.querySelector('a[href="#key-features"]');
    if (a) { a.click(); return 'link-clicked'; }
    location.hash = '#key-features';
    return 'hash-set';
}
"""

# 전체 피처 파싱 JS (두 가지 페이지 타입 모두 처리)
_PARSE_FEATURES_JS = """
(categories) => {
    const result = {};
    categories.forEach(cat => { result[cat] = {}; });

    // ── 타입 A: 신형 페이지 ──────────────────────────────────────────────
    // .model-page-features__options-list__container
    //   p "Standard Features" / "Optional Features"
    //   ul > li > button[data-category-name][data-option-name]
    const typeAContainers = [
        ...document.querySelectorAll('.model-page-features__options-list__container')
    ];
    typeAContainers.forEach(container => {
        let availability = null;
        for (const child of container.children) {
            if (child.tagName === 'P') {
                const text = child.textContent.trim();
                if (text === 'Standard Features')      availability = 'Standard';
                else if (text === 'Optional Features') availability = 'Optional';
                continue;
            }
            if (child.tagName === 'UL' && availability) {
                child.querySelectorAll('button[data-option-name]').forEach(btn => {
                    const cat  = btn.getAttribute('data-category-name');
                    const name = btn.getAttribute('data-option-name')?.trim();
                    if (cat && name && result[cat] !== undefined) {
                        result[cat][name] = availability;
                    }
                });
            }
        }
    });

    // ── 타입 B: 구형 페이지 ──────────────────────────────────────────────
    // #features-section-plugin-container
    //   .accordion-heading-text  → 카테고리명
    //   .accordion-standard-block → Standard 피처
    //   .accordion-optional-block → Optional 피처
    //   .accordion-item-cell .animate-underline--text → 피처명
    const totalA = Object.values(result).reduce((s, v) => s + Object.keys(v).length, 0);
    if (totalA === 0) {
        const container = document.getElementById('features-section-plugin-container');
        if (container) {
            container.querySelectorAll('.accordion-item').forEach(item => {
                const catEl = item.querySelector('.accordion-heading-text');
                if (!catEl) return;
                const cat = catEl.textContent.trim();
                if (!result[cat]) return;

                item.querySelectorAll('.accordion-standard-block .accordion-item-cell').forEach(cell => {
                    const name = cell.querySelector('.animate-underline--text')?.textContent.trim();
                    if (name) result[cat][name] = 'Standard';
                });
                item.querySelectorAll('.accordion-optional-block .accordion-item-cell').forEach(cell => {
                    const name = cell.querySelector('.animate-underline--text')?.textContent.trim();
                    if (name) result[cat][name] = 'Optional';
                });
            });
        }
    }

    return result;
}
"""


class BenzFeaturesScraper:

    def __init__(self, config: dict, headless: bool = True):
        self.config = config
        self.headless = headless
        self.brand = config["brand"]
        self.categories = config.get("categories", CATEGORIES)
        self.wait_strategy = config.get("wait_strategy", {})

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def scrape_model(self, model: str) -> dict | None:
        model_config = self.config["models"].get(model)
        if not model_config:
            print(f"[WARN] Model config not found: {model}")
            return None

        trim_entries = model_config.get("trims", [])
        if not trim_entries:
            print(f"[WARN] No trim URLs for model: {model}")
            return None

        all_trims: dict = {}

        with sync_playwright() as p:
            browser = p.chromium.launch(
                headless=self.headless,
                args=["--disable-http2"],
            )
            page = browser.new_page(
                viewport={"width": 1920, "height": 1080},
                user_agent=(
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/124.0.0.0 Safari/537.36"
                ),
            )

            for entry in trim_entries:
                url = entry["url"]
                print(f"\n[INFO] ===== Scraping: {url}")
                trim_data = self._scrape_trim(page, url)
                if trim_data is None:
                    print(f"[WARN] Failed — skip {url}")
                    continue
                trim_name = trim_data["trim_name"]
                all_trims[trim_name] = trim_data
                print(f"[INFO] Done — {trim_name}")

            browser.close()

        if not all_trims:
            return None

        result = self._compute_availability(all_trims)
        result.update({
            "brand": self.brand,
            "model": model,
            "crawled_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        })
        return result

    def save_excel(
        self,
        all_model_data: list[dict],
        output_dir: str = "storage/csv",
        filename: str = "benz_features.xlsx",
    ) -> str:
        """
        여러 모델 데이터를 하나의 Excel 파일에 저장.
        모델별로 시트 하나씩 생성.

        all_model_data: [scrape_model() 반환값, ...]
        """
        Path(output_dir).mkdir(parents=True, exist_ok=True)
        filepath = f"{output_dir}/{filename}"

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # 기본 Sheet 제거

        # 스타일 정의
        header_font  = Font(bold=True, color="FFFFFF")
        header_fill  = PatternFill("solid", fgColor="203864")   # 진한 파란색
        cat_font     = Font(bold=True)
        cat_fill     = PatternFill("solid", fgColor="D9E1F2")   # 연한 파란색

        VALUE_COLORS = {
            "Standard":    "E2EFDA",  # 연초록
            "Optional":    "FFF2CC",  # 연노랑
            "Unavailable": "F4CCCC",  # 연빨강
        }

        for data in all_model_data:
            if data is None:
                continue

            model     = data["model"].upper()
            trim_names = data["trim_names"]
            rows       = data["rows"]

            # 시트명: 31자 제한, 특수문자 제거
            sheet_name = model[:31].replace("/", "-").replace("\\", "-").replace("*", "").replace("?", "").replace("[", "").replace("]", "")
            ws = wb.create_sheet(title=sheet_name)

            # ── 헤더 행 ──────────────────────────────────────────────
            header = ["Category", "Feature"] + trim_names
            for col_idx, val in enumerate(header, 1):
                cell = ws.cell(row=1, column=col_idx, value=val)
                cell.font      = header_font
                cell.fill      = header_fill
                cell.alignment = Alignment(horizontal="center", wrap_text=True)

            # ── 데이터 행 ─────────────────────────────────────────────
            for row_idx, row in enumerate(rows, 2):
                cat     = row["category"]
                feature = row["feature"]

                # Category 열
                cat_cell = ws.cell(row=row_idx, column=1, value=cat)
                cat_cell.font = cat_font
                cat_cell.fill = cat_fill

                # Feature 열
                ws.cell(row=row_idx, column=2, value=feature)

                # 트림별 값
                for col_idx, trim_name in enumerate(trim_names, 3):
                    val  = row["values"].get(trim_name, "Unavailable")
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    color = VALUE_COLORS.get(val)
                    if color:
                        cell.fill = PatternFill("solid", fgColor=color)
                    cell.alignment = Alignment(horizontal="center")

            # ── 열 너비 자동 조정 ──────────────────────────────────────
            ws.column_dimensions["A"].width = 14   # Category
            ws.column_dimensions["B"].width = 50   # Feature
            trim_col_width = max(18, min(30, 80 // max(len(trim_names), 1)))
            for col_idx in range(3, 3 + len(trim_names)):
                col_letter = openpyxl.utils.get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = trim_col_width

            # 첫 행 고정
            ws.freeze_panes = "C2"

            print(f"[INFO] Sheet written: {sheet_name} ({len(rows)} rows, {len(trim_names)} trims)")

        wb.save(filepath)
        print(f"[INFO] Excel saved → {filepath}")
        return filepath

    def save_csv(self, data: dict, output_dir: str = "storage/csv") -> str:
        """단일 모델 결과를 CSV로 저장 (보조 용도)."""
        Path(output_dir).mkdir(parents=True, exist_ok=True)
        model_slug = data["model"].replace(" ", "_").lower()
        filename = f"{output_dir}/benz_features_{model_slug}.csv"

        trim_names = data["trim_names"]
        rows = data["rows"]

        with open(filename, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(["Category", "Feature"] + trim_names)
            for row in rows:
                writer.writerow(
                    [row["category"], row["feature"]]
                    + [row["values"].get(t, "Unavailable") for t in trim_names]
                )

        print(f"[INFO] CSV saved → {filename}")
        return filename

    # ------------------------------------------------------------------
    # Single trim
    # ------------------------------------------------------------------

    def _scrape_trim(self, page, url: str) -> dict | None:
        try:
            wait_until = self.wait_strategy.get("wait_until", "domcontentloaded")
            page.goto(url, wait_until=wait_until, timeout=60000)
        except PlaywrightTimeoutError:
            print(f"[ERROR] Timeout: {url}")
            return None

        time.sleep(1.5)

        # usercentrics 오버레이 제거 (JS로 직접 처리)
        consent = page.evaluate(_CONSENT_JS)
        print(f"[INFO] Consent: {consent}")
        if consent == "shadow-clicked":
            time.sleep(1.0)

        # 트림명 추출
        trim_name = self._get_trim_name(page, url)
        print(f"[INFO] Trim name: {trim_name}")

        # 스크롤하여 sticky nav lazy rendering 트리거
        page.evaluate("window.scrollBy(0, 600)")
        time.sleep(1.0)
        page.evaluate("window.scrollTo(0, 0)")
        time.sleep(0.5)

        # sticky nav의 Key Features 링크 대기
        try:
            page.wait_for_selector(
                "a[href='#key-features']",
                timeout=15000,
                state="attached",
            )
            print("[INFO] sticky nav ready")
        except PlaywrightTimeoutError:
            # 없으면 더 깊이 스크롤 후 재시도
            print("[INFO] sticky nav not yet ready — scrolling further")
            page.evaluate("window.scrollBy(0, 1200)")
            time.sleep(1.5)
            page.evaluate("window.scrollTo(0, 0)")
            time.sleep(0.5)

        # Key Features 탭 활성화 (JS click — overlay 우회)
        kf = page.evaluate(_ACTIVATE_KF_JS)
        print(f"[INFO] Key Features activation: {kf}")
        time.sleep(1.5)

        # 피처 데이터 DOM 삽입 대기
        if not self._wait_for_features(page):
            # hash 방식 실패 시: 페이지를 anchor URL로 직접 재진입
            print("[INFO] Trying direct anchor URL navigation")
            try:
                page.goto(url + "#key-features", wait_until="domcontentloaded", timeout=30000)
            except PlaywrightTimeoutError:
                pass
            time.sleep(2.0)
            page.evaluate(_CONSENT_JS)
            if not self._wait_for_features(page):
                print("[WARN] Feature list not found — saving debug snapshot")
                slug = trim_name.replace(" ", "_")[:30]
                self._save_debug(page, f"debug_feat_{slug}")
                return None

        # 전체 피처 일괄 파싱
        features_by_category: dict[str, dict[str, str]] = page.evaluate(
            _PARSE_FEATURES_JS, self.categories
        )

        total = sum(len(v) for v in features_by_category.values())
        print(f"[INFO] Features extracted: {total}")
        for cat, feats in features_by_category.items():
            if feats:
                print(f"   {cat}: {len(feats)}")

        if total == 0:
            print("[WARN] 0 features — saving debug snapshot")
            self._save_debug(page, f"debug_feat_{trim_name.replace(' ', '_')[:30]}")

        return {
            "trim_name": trim_name,
            "url": url,
            "features_by_category": features_by_category,
        }

    # ------------------------------------------------------------------
    # Availability computation
    # ------------------------------------------------------------------

    def _compute_availability(self, all_trims: dict) -> dict:
        """
        전 트림 비교:
          Standard   → 해당 트림에서 기본 포함
          Optional   → 해당 트림에서 옵션
          Unavailable → 다른 트림에는 있으나 이 트림에 없음
        """
        trim_names = list(all_trims.keys())
        # dict를 순서 보존 집합으로 사용 (첫 번째 트림 기준 순서 유지)
        universe: dict[str, dict] = {cat: {} for cat in self.categories}

        for trim_data in all_trims.values():
            for cat in self.categories:
                for feat in trim_data["features_by_category"].get(cat, {}).keys():
                    universe[cat].setdefault(feat, None)

        rows = []
        for cat in self.categories:
            for feature in universe[cat]:
                values = {}
                for trim_name in trim_names:
                    cat_features = all_trims[trim_name]["features_by_category"].get(cat, {})
                    values[trim_name] = cat_features.get(feature, "Unavailable")
                rows.append({"category": cat, "feature": feature, "values": values})

        return {"trim_names": trim_names, "rows": rows}

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    def _wait_for_features(self, page, timeout: int = 25000) -> bool:
        """피처 목록 DOM 삽입 대기 (타입 A/B 모두). 성공 시 True."""
        # 타입 A: 신형 (Key Features 탭 클릭 후 삽입)
        # 타입 B: 구형 (features-section-plugin-container, 초기 DOM에 존재)
        selector = (
            "p.model-page-features__options-list__title, "
            "#features-section-plugin-container .accordion-item-cell"
        )
        try:
            page.wait_for_selector(selector, timeout=timeout, state="attached")
            print("[INFO] Feature list loaded")
            return True
        except PlaywrightTimeoutError:
            return False

    def _get_trim_name(self, page, url: str) -> str:
        """네비게이션 h1 제외 후 실제 차량명 추출."""
        EXCLUDE = {"Vehicles", "Models", "Electric", "AMG", ""}

        # 특정 클래스 우선 시도
        for sel in [
            ".vehicle-hero__name",
            ".vehicle-hero__title",
            ".pdp-hero__name",
            ".model-page-hero__name",
        ]:
            try:
                el = page.query_selector(sel)
                if el:
                    text = el.text_content().replace("  ", " ").strip()
                    if text and len(text) > 3:
                        return text
            except Exception:
                pass

        # nav/header 밖의 h1 중 의미있는 것
        try:
            name = page.evaluate(
                """
                (exclude) => {
                    for (const h of document.querySelectorAll('h1')) {
                        const t = h.textContent.replace(/\\s+/g, ' ').trim();
                        if (!t || exclude.includes(t)) continue;
                        if (h.closest('nav, header, .main-navigation, .mobile-nav')) continue;
                        return t;
                    }
                    return null;
                }
                """,
                list(EXCLUDE),
            )
            if name:
                return name
        except Exception:
            pass

        return url.rstrip("/").split("/")[-1]

    def _save_debug(self, page, prefix: str) -> None:
        try:
            page.screenshot(path=f"{prefix}.png", full_page=False)
            with open(f"{prefix}.html", "w", encoding="utf-8") as f:
                f.write(page.content())
            print(f"[DEBUG] Saved: {prefix}.png / .html")
        except Exception as e:
            print(f"[WARN] Debug save failed: {e}")
